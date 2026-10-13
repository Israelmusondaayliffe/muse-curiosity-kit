"""Offline, explicitly exported review packets. Standard library, Python 3.9+."""
from __future__ import annotations
import argparse
import base64
import hashlib
import json
import math
import os
import re
import tempfile
import io
import zipfile
from pathlib import Path, PurePosixPath
from .common import HectlError

SCHEMA = "he-review-v1"
FEEDBACK = "he-feedback-v1"
MAX_JSON = 4 * 1024 * 1024
MAX_FILE = 8 * 1024 * 1024
MAX_TOTAL = 24 * 1024 * 1024
ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,127}$")
SHA = re.compile(r"^[0-9a-f]{64}$")
TEXT = {".txt", ".md", ".json", ".csv", ".tsv", ".py", ".js", ".css", ".html", ".svg", ".log", ".yaml", ".yml", ".xml"}

def canonical(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False).encode("utf-8")

def digest(value):
    return hashlib.sha256(canonical(value)).hexdigest()

def _text(value, name, maximum=200000, nonempty=False):
    if not isinstance(value, str) or len(value) > maximum or (nonempty and not value.strip()):
        raise HectlError("invalid " + name)
    return value

def _id(value, name):
    if not isinstance(value, str) or not ID.fullmatch(value):
        raise HectlError("invalid " + name)
    return value

def _sha(value, name):
    if not isinstance(value, str) or not SHA.fullmatch(value):
        raise HectlError("invalid " + name)
    return value

def _finite(value):
    if isinstance(value, float) and not math.isfinite(value):
        raise HectlError("non-finite number")
    if isinstance(value, dict):
        for k, v in value.items():
            _text(k, "object key", 256)
            _finite(v)
    elif isinstance(value, list):
        for v in value:
            _finite(v)

def load(path):
    p = Path(path)
    if p.is_symlink() or not p.is_file() or p.stat().st_size > MAX_JSON:
        raise HectlError("review JSON must be a regular file within size limit")
    try:
        value = json.loads(p.read_text(encoding="utf-8"), parse_constant=lambda _: (_ for _ in ()).throw(ValueError("nonfinite")))
        _finite(value)
        return value
    except (ValueError, UnicodeError, RecursionError) as exc:
        raise HectlError("invalid review JSON: " + str(exc)) from exc

def validate_queries(queries):
    if not isinstance(queries, list) or not 1 <= len(queries) <= 1000:
        raise HectlError("queries must contain 1 to 1000 cases")
    ids, seen = set(), {}
    for q in queries:
        if not isinstance(q, dict) or set(q) - {"case_id", "query", "should_trigger", "partition"}:
            raise HectlError("invalid query fields")
        cid = _id(q.get("case_id"), "query case_id")
        if cid in ids:
            raise HectlError("duplicate query case_id")
        ids.add(cid)
        query = _text(q.get("query"), "query", 20000, True)
        if type(q.get("should_trigger")) is not bool:
            raise HectlError("should_trigger must be boolean")
        partition = q.get("partition", "development")
        if partition not in ("development", "validation", "final"):
            raise HectlError("invalid query partition")
        if query in seen and seen[query] != (q["should_trigger"], partition):
            raise HectlError("duplicate query has conflicting label or partition")
        seen[query] = (q["should_trigger"], partition)
    return queries

def validate_manifest(manifest):
    if not isinstance(manifest, dict) or manifest.get("schema") != SCHEMA:
        raise HectlError("expected " + SCHEMA)
    if set(manifest) - {"schema", "mode", "title", "queries", "history", "runs", "benchmark", "notes"}:
        raise HectlError("unknown review manifest fields")
    mode = manifest.get("mode")
    if mode not in ("trigger", "history", "artifacts"):
        raise HectlError("mode must be trigger, history, or artifacts")
    _text(manifest.get("title", "Review"), "title", 300)
    _text(manifest.get("notes", ""), "notes")
    if mode == "trigger":
        validate_queries(manifest.get("queries"))
    if mode == "history":
        history = manifest.get("history")
        if not isinstance(history, list) or len(history) > 1000:
            raise HectlError("history must be a list with at most 1000 entries")
        ids = set()
        for item in history:
            if not isinstance(item, dict) or set(item) - {"iteration_id", "description", "metrics", "selection_rationale", "partition", "status"}:
                raise HectlError("invalid history entry")
            ident = _id(item.get("iteration_id"), "iteration_id")
            if ident in ids:
                raise HectlError("duplicate iteration_id")
            ids.add(ident)
            _text(item.get("description"), "description")
            _text(item.get("selection_rationale", ""), "selection_rationale")
            _text(item.get("status", "unknown"), "history status", 128)
            _text(item.get("partition", "unknown"), "history partition", 128)
    if mode == "artifacts":
        runs = manifest.get("runs")
        if not isinstance(runs, list) or not 1 <= len(runs) <= 1000:
            raise HectlError("runs must contain 1 to 1000 entries")
        ids = set()
        for run in runs:
            if not isinstance(run, dict) or set(run) - {"run_id", "case_id", "case_digest", "run_digest", "target_digest", "role", "label", "artifacts", "grading", "metrics", "notes", "status"}:
                raise HectlError("invalid run fields")
            rid = _id(run.get("run_id"), "run_id")
            if (rid, run.get("case_id")) in ids:
                raise HectlError("duplicate run_id/case_id")
            ids.add((rid, run.get("case_id")))
            _id(run.get("case_id"), "case_id")
            _sha(run.get("case_digest"), "case_digest")
            _sha(run.get("run_digest"), "run_digest")
            _sha(run.get("target_digest"), "target_digest")
            if run.get("role") not in ("current", "prior"):
                raise HectlError("run role must be current or prior")
            for field in ("label", "notes", "status"):
                _text(run.get(field, ""), field)
            arts = run.get("artifacts")
            if not isinstance(arts, list) or len(arts) > 200:
                raise HectlError("artifacts must be an explicit list, at most 200 per run")
            paths = set()
            for art in arts:
                if not isinstance(art, dict) or set(art) != {"path", "sha256"}:
                    raise HectlError("artifact requires exactly path and sha256")
                p = _text(art["path"], "artifact path", 1000, True)
                parts = PurePosixPath(p).parts
                if "\\" in p or p.startswith("/") or any(x in ("", ".", "..") for x in p.split("/")) or not parts:
                    raise HectlError("artifact path must be a contained relative POSIX path")
                if p in paths:
                    raise HectlError("duplicate artifact path")
                paths.add(p)
                _sha(art["sha256"], "artifact sha256")
    _finite(manifest)
    if len(canonical(manifest)) > MAX_JSON:
        raise HectlError("manifest exceeds size cap")
    return manifest

def _workbook_preview(content):
    """Optional installed parser; formulas are displayed, never executed."""
    try:
        import openpyxl
    except ImportError:
        return None, "Download only. Trusted openpyxl parser unavailable; all-sheet review is incomplete."
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(x.file_size for x in members) > 32 * 1024 * 1024:
                return None, "Workbook expanded content exceeds safe preview limits; all-sheet review is incomplete."
            for member in members:
                if member.filename.lower().endswith((".xml", ".rels")):
                    xml = archive.read(member)
                    if b"\x00" in xml or b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        return None, "Workbook XML declarations or encoding are unsafe for preview; all-sheet review is incomplete."
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheets, total, chars, complete = [], 0, 0, True
        try:
            if len(book.worksheets) > 50:
                return None, "Workbook exceeds 50-sheet limit; all-sheet review is incomplete."
            for sheet in book.worksheets:
                rows = []
                limited = (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 100
                complete = complete and not limited
                for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, 500), max_col=min(sheet.max_column or 0, 100), values_only=True):
                    if total + len(row) > 50000:
                        limited, complete = True, False
                        break
                    raw_values = [None if v is None else str(v) for v in row]
                    if any(v is not None and len(v) > 20000 for v in raw_values):
                        limited, complete = True, False
                    values = [None if v is None else v[:20000] for v in raw_values]
                    chars += sum(len(v) for v in values if v is not None)
                    if chars > 2 * 1024 * 1024:
                        limited, complete = True, False
                        break
                    rows.append(values)
                    total += len(row)
                sheets.append({"name": sheet.title, "visibility": sheet.sheet_state, "rows": rows, "truncated": limited})
        finally:
            book.close()
        limitation = "Cell-value and formula-text preview only; formulas are not recalculated. Charts, formatting, images and print layout require a trusted workbook application. "
        limitation += "All worksheet tabs included." if complete else "Preview truncated; required all-sheet/cell review is incomplete."
        return {"sheets": sheets, "complete_cells": complete}, limitation
    except Exception as exc:
        return None, "Workbook preview failed (" + exc.__class__.__name__ + "); download only, all-sheet review incomplete."

def _artifact(root, item):
    path = root
    for part in PurePosixPath(item["path"]).parts:
        path = path / part
        if path.is_symlink():
            raise HectlError("symlink artifact refused: " + item["path"])
    try:
        path.resolve().relative_to(root.resolve())
    except ValueError as exc:
        raise HectlError("artifact outside evidence root") from exc
    result = dict(item)
    if not path.is_file():
        return dict(result, kind="error", error="Missing or unreadable regular artifact")
    size = path.stat().st_size
    if size > MAX_FILE:
        return dict(result, kind="error", error="Artifact exceeds 8 MiB preview/export cap", size=size)
    try:
        # Cap the actual read too, including files changing during capture.
        with path.open("rb") as stream:
            content = stream.read(MAX_FILE + 1)
    except OSError as exc:
        return dict(result, kind="error", error="Artifact could not be read: " + exc.__class__.__name__)
    if len(content) > MAX_FILE or hashlib.sha256(content).hexdigest() != item["sha256"]:
        raise HectlError("artifact changed or digest mismatch: " + item["path"])
    suffix = path.suffix.lower()
    kind, mime, text = "binary", "application/octet-stream", None
    if suffix == ".png" and content.startswith(b"\x89PNG\r\n\x1a\n"):
        kind, mime = "image", "image/png"
    elif suffix in (".jpg", ".jpeg") and content.startswith(b"\xff\xd8\xff"):
        kind, mime = "image", "image/jpeg"
    elif suffix == ".gif" and content[:6] in (b"GIF87a", b"GIF89a"):
        kind, mime = "image", "image/gif"
    elif suffix == ".webp" and content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        kind, mime = "image", "image/webp"
    elif suffix in TEXT:
        try:
            text = content.decode("utf-8")
            kind = "text"
        except UnicodeError:
            pass
    result.update(kind=kind, mime=mime, size=len(content), data=base64.b64encode(content).decode("ascii"))
    if text is not None:
        result["text"] = text
    if suffix == ".xlsx":
        workbook, limitation = _workbook_preview(content)
        result["limitation"] = limitation
        if workbook is not None:
            result.update(kind="workbook", workbook=workbook)
    elif suffix in (".xls", ".ods"):
        result["limitation"] = "Download only. No supported trusted parser for this format; required all-sheet review remains incomplete."
    elif suffix == ".pdf":
        result["limitation"] = "Download only. PDF content is not executed or embedded; required visual page review remains incomplete. Open only in a trusted viewer."
    elif suffix in (".html", ".svg", ".js"):
        result["limitation"] = "Inert source preview only. Downloaded active content must be treated as untrusted."
    return result

def identities(manifest):
    answer = {}
    for run in manifest.get("runs", []):
        base = {k: run[k] for k in ("case_id", "case_digest", "run_id", "run_digest", "target_digest", "role")}
        for art in run["artifacts"]:
            identity = dict(base, artifact_path=art["path"], artifact_sha256=art["sha256"])
            answer[digest(identity)] = identity
    return answer

def validate_feedback(feedback, manifest):
    validate_manifest(manifest)
    if not isinstance(feedback, dict) or set(feedback) != {"schema", "packet_digest", "queries", "reviews"}:
        raise HectlError("invalid feedback fields")
    if feedback["schema"] != FEEDBACK or feedback["packet_digest"] != digest(manifest):
        raise HectlError("feedback belongs to a different review packet")
    if manifest["mode"] == "trigger":
        validate_queries(feedback["queries"])
    elif feedback["queries"] is not None:
        raise HectlError("only trigger mode can change queries")
    reviews = feedback["reviews"]
    if not isinstance(reviews, dict):
        raise HectlError("reviews must be an object")
    allowed = identities(manifest)
    for key, row in reviews.items():
        if key not in allowed or not isinstance(row, dict) or set(row) != {"identity", "reviewed", "accepted", "comment", "reviewer", "evidence"}:
            raise HectlError("invalid review identity or fields")
        if row["identity"] != allowed[key]:
            raise HectlError("review identity mismatch")
        if type(row["reviewed"]) is not bool or row["accepted"] not in (None, True, False) or (row["accepted"] is not None and type(row["accepted"]) is not bool):
            raise HectlError("invalid reviewed/accepted state")
        if row["accepted"] is not None and not row["reviewed"]:
            raise HectlError("accept/reject requires explicit reviewed state")
        _text(row["comment"], "review comment", 20000)
        _text(row["reviewer"], "reviewer", 300, row["accepted"] is not None)
        _text(row["evidence"], "review evidence", 20000, row["accepted"] is not None)
    return feedback

def _write_new(path, data):
    path = Path(path)
    if path.exists() or path.is_symlink():
        raise HectlError("output already exists; choose a new path")
    if any(p.is_symlink() for p in [path.parent] + list(path.parent.parents)):
        raise HectlError("symlink output parent refused")
    if not path.parent.is_dir():
        raise HectlError("output parent must exist")
    fd, temp = tempfile.mkstemp(prefix=".he-review-", dir=str(path.parent))
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        # Hard-link gives atomic create without overwriting a racing existing path.
        os.link(temp, path)
    finally:
        os.unlink(temp)

def build(manifest, evidence_root, output, feedback=None):
    validate_manifest(manifest)
    root = Path(evidence_root)
    if any(p.is_symlink() for p in [root] + list(root.parents)) or not root.is_dir():
        raise HectlError("evidence root must be an existing nonsymlink directory")
    output = Path(output)
    try:
        output.resolve().relative_to(root.resolve())
    except ValueError:
        pass
    else:
        raise HectlError("review output must be outside evidence root")
    packet = {"manifest": manifest, "packet_digest": digest(manifest), "identities": identities(manifest), "artifacts": {}, "feedback": None}
    total = 0
    for key, identity in packet["identities"].items():
        art = _artifact(root, {"path": identity["artifact_path"], "sha256": identity["artifact_sha256"]})
        total += art.get("size", 0) if art["kind"] != "error" else 0
        if total > MAX_TOTAL:
            raise HectlError("selected artifacts exceed 24 MiB total export cap")
        packet["artifacts"][key] = art
    if feedback is not None:
        packet["feedback"] = validate_feedback(feedback, manifest)
    template = Path(__file__).resolve().parents[2] / "assets" / "review.html"
    html = template.read_text(encoding="utf-8")
    payload = base64.b64encode(canonical(packet)).decode("ascii")
    html = html.replace("__HE_PACKET_BASE64__", payload)
    scripts = re.findall(r"<script>(.*?)</script>", html, flags=re.S)
    if len(scripts) != 1:
        raise HectlError("review template requires one inline script")
    script_hash = base64.b64encode(hashlib.sha256(scripts[0].encode()).digest()).decode()
    html = html.replace("__HE_SCRIPT_HASH__", script_hash)
    _write_new(output, html.encode("utf-8"))
    return {"status": "prepared", "packet_digest": packet["packet_digest"], "output": str(output), "artifact_count": len(packet["artifacts"]), "limits": "Offline export. Imported feedback is declared human review, not authenticated consent or release acceptance."}

def dispatch(argv):
    parser = argparse.ArgumentParser(prog="hectl review", description="Prepare offline review or validate imported review feedback.")
    sub = parser.add_subparsers(dest="action", required=True)
    p = sub.add_parser("build")
    p.add_argument("--manifest", required=True)
    p.add_argument("--evidence-root", required=True)
    p.add_argument("--output", required=True)
    p.add_argument("--feedback")
    p = sub.add_parser("import")
    p.add_argument("--manifest", required=True)
    p.add_argument("--feedback", required=True)
    p.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    try:
        manifest = load(args.manifest)
        if args.action == "build":
            result = build(manifest, args.evidence_root, args.output, load(args.feedback) if args.feedback else None)
        else:
            feedback = validate_feedback(load(args.feedback), manifest)
            _write_new(Path(args.output), canonical(feedback) + b"\n")
            result = {"status": "imported-declared-review", "packet_digest": digest(manifest), "reviews": len(feedback["reviews"]), "output": args.output}
    except (HectlError, OSError, ValueError, KeyError, TypeError) as exc:
        print(json.dumps({"error": str(exc)}))
        return 1
    print(json.dumps(result, indent=2))
    return 0
